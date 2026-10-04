"""Builds Heart_Disease_Presentation.pptx from results.json + figures/ (run the pipeline first)."""
import json
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

BASE = Path(__file__).parent
R = json.loads((BASE / "results.json").read_text())
FIG = BASE / "figures"
NAVY, RED, GREY = RGBColor(0x1F, 0x2A, 0x44), RGBColor(0xD1, 0x49, 0x5B), RGBColor(0x55, 0x55, 0x55)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
blank = prs.slide_layouts[6]


def text(slide, s, x, y, w, h, size=18, bold=False, color=GREY):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.space_after = Pt(8)
        p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, color
    return tb


def slide(title):
    s = prs.slides.add_slide(blank)
    text(s, title, 0.7, 0.4, 12, 0.9, 32, True, NAVY)
    bar = s.shapes.add_shape(1, Inches(0.7), Inches(1.3), Inches(1.2), Inches(0.07))
    bar.fill.solid(); bar.fill.fore_color.rgb = RED; bar.line.fill.background()
    return s


def pic(s, name, x, y, w):
    if (FIG / name).exists():
        s.shapes.add_picture(str(FIG / name), Inches(x), Inches(y), width=Inches(w))


M, c = R["models"], R["cleaning"]
best = R["selected_model"]; b = M[best]

s = prs.slides.add_slide(blank)
text(s, "Heart Disease Prediction", 0.9, 2.4, 11.5, 1.2, 48, True, NAVY)
text(s, "Machine Learning & AI - Week 4 Capstone Project", 0.9, 3.7, 11, 0.6, 22)
text(s, "Binary classification with Logistic Regression and Decision Tree", 0.9, 4.3, 11, 0.6, 18)

s = slide("Problem & Dataset")
text(s, ["Goal: predict whether a patient is likely to have heart disease from clinical measurements.",
         "Dataset: Kaggle - Heart Disease UCI.",
         f"{c['rows_clean']} patients after cleaning ({c['rows_raw']} raw rows, "
         f"{c['duplicates_removed']} duplicates removed).",
         f"{len(R['numeric_features'])} numeric and {len(R['categorical_features'])} categorical features.",
         "Target: 0 = no disease, 1 = disease."], 0.7, 1.7, 6.2, 5, 20)
pic(s, "class_balance.png", 7.4, 1.8, 5.2)

s = slide("Data Cleaning & Preprocessing")
miss = c["missing_values"]
text(s, ["Target converted to binary (any disease grade > 0 becomes 1).",
         "Identifier columns and duplicate rows removed.",
         "Missing values: " + (", ".join(f"{k} ({v})" for k, v in miss.items()) if miss else "none"),
         "Numeric features: median imputation, then StandardScaler.",
         "Categorical features: most-frequent imputation, then one-hot encoding.",
         "Everything sits inside one scikit-learn Pipeline, so no test data leaks into training."],
     0.7, 1.7, 6.4, 5, 19)
pic(s, "correlation.png", 7.4, 1.7, 5.3)

s = slide("Method")
text(s, ["Stratified 80/20 train-test split (random_state = 42).",
         "Models: Logistic Regression and Decision Tree.",
         "5-fold stratified cross-validation for baseline accuracy.",
         "GridSearchCV tuning (scored on ROC AUC): regularisation C for Logistic Regression; "
         "depth, leaf size and criterion for the tree.",
         "Final metrics computed once on the held-out test set."], 0.7, 1.7, 12, 5, 21)

s = slide("Results")
rows = [("Model", "CV acc.", "Accuracy", "Precision", "Recall", "F1", "ROC AUC")]
for n, m in M.items():
    rows.append((n, *(f"{m[k]:.3f}" for k in ("cv_accuracy_mean", "accuracy", "precision", "recall", "f1", "roc_auc"))))
tbl = s.shapes.add_table(len(rows), 7, Inches(0.7), Inches(1.8), Inches(11.9), Inches(0.6 * len(rows))).table
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        cell = tbl.cell(i, j); cell.text = v
        cell.text_frame.paragraphs[0].font.size = Pt(18)
        cell.text_frame.paragraphs[0].font.bold = (i == 0)
text(s, f"Selected model: {best} (highest ROC AUC, with F1 as tie-breaker).",
     0.7, 4.3, 12, 1, 20, True, NAVY)

s = slide("Confusion Matrices")
x = 0.7
for n in M:
    pic(s, f"confusion_{n.lower().replace(' ', '_')}.png", x, 1.7, 5.6); x += 6.2

s = slide("ROC Curve")
pic(s, "roc_curve.png", 0.7, 1.6, 6.2)
text(s, [f"{n}: AUC = {m['roc_auc']:.3f}" for n, m in M.items()] +
     ["A curve closer to the top-left corner means better separation between sick and healthy patients."],
     7.4, 2.2, 5.3, 4, 20)

s = slide("Conclusion")
tn, fp = b["confusion_matrix"][0]; fn, tp = b["confusion_matrix"][1]
text(s, [f"{best} performed best: ROC AUC {b['roc_auc']:.3f}, accuracy {b['accuracy']:.1%}, recall {b['recall']:.1%}, precision {b['precision']:.1%}.",
         f"On {R['test_size']} unseen patients it caught {tp} of {tp + fn} true disease cases and missed {fn}.",
         "The model is saved with joblib (models/heart_model.joblib) and can predict on new patients.",
         "Limitations: small dataset, single train-test split; this is a learning project, not a diagnostic tool.",
         "Next steps: try Random Forest, add calibration, test on external data."], 0.7, 1.7, 12, 5, 21)

prs.save(BASE / "Heart_Disease_Presentation.pptx")
print("Saved Heart_Disease_Presentation.pptx")
